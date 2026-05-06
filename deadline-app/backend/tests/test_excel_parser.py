import pytest
from datetime import date, datetime
from services.excel_parser import ExcelParser
import io

# -- helpers ---------------------------------

def make_parser():
    # create a workbook to instantiate
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Theo dõi CV Tháng 5"
    wb.create_sheet("Chỉ đạo LĐP giao ban")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return ExcelParser(buf)

# -- extract_deadline-------------------------

class TestExtractDeadline:

    def setup_method(self):
        self.parser = make_parser()

    def test_extract_date_from_free_text(self):
        result = self.parser.extract_deadline("KS. Phú hoàn thành trước 15/5/2026")
        assert result == date(2026, 5, 15)

    def test_extract_last_date_when_multiple(self):
        #from X to Y - should return the deadline Y
        result = self.parser.extract_deadline("từ 1/5/2026 đến 15/5/2026")
        assert result == date(2026, 5, 15)

    def test_extract_date_object_directly(self):
        result = self.parser.extract_deadline(date(2026, 5, 15))
        assert result == date(2026, 5, 15)

    def test_returns_none_for_empty(self):
        assert self.parser.extract_deadline("") is None
        assert self.parser.extract_deadline(None) is None

    def test_returns_none_for_no_date(self):
        result = self.parser.extract_deadline("Lưu hồ sơ theo dõi")
        assert result is None

    def test_single_digit_day_month(self):
        result = self.parser.extract_deadline("hoàn thành trước 1/6/2026")
        assert result == date(2026, 6, 1)

    def test_extract_datetime_object(self):
        result = self.parser.extract_deadline(datetime(2026, 5, 15, 9, 30))
        assert result == date(2026, 5, 15)

# -- detect_recurring-------------------------

class TestDetectRecurring:

    def setup_method(self):
        self.parser = make_parser()

    def test_everyday(self):
        is_recurring, label = self.parser.detect_recurring("Mỗi ngày")
        assert is_recurring is True
        assert label == "Mỗi ngày"

    def test_everyday_2(self):
        is_recurring, label = self.parser.detect_recurring("Mỗi ngày")
        assert is_recurring is True
        assert label == "Mỗi ngày"

    def test_weekly(self):
        is_recurring, _ = self.parser.detect_recurring("hàng tuần vào thứ 2")
        assert is_recurring is True
  
    def test_monthly(self):
        is_recurring, _ = self.parser.detect_recurring("Hàng tháng")
        assert is_recurring is True

    def test_frequently(self):
        is_recurring, _ = self.parser.detect_recurring("thường xuyên kiểm tra")
        assert is_recurring is True

    def test_not_recurring(self):
        is_recurring, label = self.parser.detect_recurring("hoàn thành trước 15/5/2026")
        assert is_recurring is False
        assert label is None

    def test_empty(self):
        is_recurring, label = self.parser.detect_recurring("")
        assert is_recurring is False
        assert label is None

    def test_none(self):
        is_recurring, label = self.parser.detect_recurring(None)
        assert is_recurring is False
        assert label is None

    def test_case_insensitive(self):
        is_recurring, _ = self.parser.detect_recurring("MỖI NGÀY")
        assert is_recurring is True

# -- extract_staff_names------------------------

class TestExtractStaffNames:

    def setup_method(self):
        self.parser = make_parser()

    def test_newline_separator(self):
        result = self.parser.extract_staff_names("KS. Tiến\nKS. Bảo")
        assert result == ["KS. Tiến", "KS. Bảo"]

    def test_comma_separator(self):
        result = self.parser.extract_staff_names("KS. Tiến, KS. Bảo")
        assert result == ["KS. Tiến", "KS. Bảo"]

    def test_semicolon_separator(self):
        result = self.parser.extract_staff_names("KS. Tiến; KS. Bảo")
        assert result == ["KS. Tiến", "KS. Bảo"]

    def test_strips_whitespace(self):
        result = self.parser.extract_staff_names("  KS. Tiến  ,  KS. Bảo  ")
        assert result == ["KS. Tiến", "KS. Bảo"]

    def test_single_name(self):
        result = self.parser.extract_staff_names("Đại Phú")
        assert result == ["Đại Phú"]

    def test_empty(self):
        assert self.parser.extract_staff_names("") == []
        assert self.parser.extract_staff_names(None) == []

    def test_multiple_newlines(self):
        result = self.parser.extract_staff_names("KS. Tiến\n\nKS. Bảo")
        assert result == ["KS. Tiến", "KS. Bảo"]

    def test_mixed_separator(self):
        result = self.parser.extract_staff_names("KS. Tiến\nKS. Bảo, KS. An")
        assert result == ["KS. Tiến", "KS. Bảo", "KS. An"]

# -- suggest_tab-------------------------------

class TestSuggestTab:

    def setup_method(self):
        self.parser = make_parser()

    def test_suggests_correct_tab1(self):
        result = self.parser._suggest_tab("theo doi cv")
        assert result == "Theo dõi CV Tháng 5"

    def test_suggests_correct_tab2(self):
        result = self.parser._suggest_tab("chi dao ldp")
        assert result == "Chỉ đạo LĐP giao ban"

    def test_monthly_name_change(self):
        # simulate next month - should still match
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.title = "Theo dõi CV Tháng 6"
        wb.create_sheet("Chỉ đạo LĐP giao ban")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        parser = ExcelParser(buf)
        result = parser._suggest_tab("theo doi cv")
        assert result == "Theo dõi CV Tháng 6"