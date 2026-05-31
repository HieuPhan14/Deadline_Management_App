import openpyxl
import re
from rapidfuzz import fuzz
from dateutil.parser import parse as parse_date
from dateutil.parser import ParserError
from datetime import date, datetime


class ExcelParser:
    TAB1_HEADER_ROW = 6
    TAB2_HEADER_ROW = 9

    def __init__(self, file):
        self.wb = openpyxl.load_workbook(file)

    def detect_tabs(self) -> dict:
        return {
            "sheet_names": self.wb.sheetnames,
            "suggested": {
                "tab1": self._suggest_tab("documents tracking theo doi cv"),
                "tab2": self._suggest_tab("directives meeting chi dao ldp")
            }
        }

    def _suggest_tab(self, keywords: str) -> str:
        def score(name):
            return fuzz.partial_ratio(keywords, name.lower())
        return max(self.wb.sheetnames, key=score)

    def extract_staff_names(self, text: str) -> list:
        if not text:
            return []
        names = re.split(r'[\n,;]+', str(text))
        return [name.strip() for name in names if name.strip() and len(name.strip()) <= 100]

    def detect_recurring(self, text: str) -> tuple:
        if not text:
            return (False, None)

        recurring_keywords = {
            "mỗi ngày", "hàng ngày", "hàng tuần",
            "hàng tháng", "thường xuyên",
            "daily", "weekly", "monthly", "yearly", "annually",
            "recurring", "recurrent", "ongoing", "regular", "regularly",
            "every day", "every week", "every month", "continuous"
        }
        text_lower = str(text).lower()

        for keyword in recurring_keywords:
            if keyword in text_lower:
                return (True, str(text).strip())

        return (False, None)

    def extract_deadline(self, text) -> date:
        if not text:
            return None

        if isinstance(text, datetime):
            return text.date()
        if isinstance(text, date):
            return text

        text = str(text)

        iso_matches = re.findall(r'\d{4}-\d{2}-\d{2}', text)
        if iso_matches:
            try:
                return parse_date(iso_matches[-1]).date()
            except ParserError:
                pass

        slash_matches = re.findall(r'\d{1,2}/\d{1,2}/\d{4}', text)
        if slash_matches:
            try:
                return parse_date(slash_matches[-1], dayfirst=True).date()
            except ParserError:
                pass

        return None

    def _parse_date(self, value) -> date:
        if value is None:
            return None
        if hasattr(value, 'date'):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return parse_date(str(value), dayfirst=True).date()
        except (ParserError, ValueError):
            return None

    def _cell_str(self, value):
        if value is None:
            return None
        return str(value).strip() or None

    def get_tab_preview(self, tab_name: str, header_row: int) -> dict:
        sheet = self.wb[tab_name]
        rows = []
        for row in sheet.iter_rows(min_row=max(1, header_row - 1), max_row=header_row + 2, values_only=True):
            rows.append(row)

        if not rows:
            return {"columns": []}

        num_cols = max(len(r) for r in rows)
        label_row = rows[0] if rows else []
        data_rows = rows[1:]

        columns = []
        for col_idx in range(num_cols):
            header_val = label_row[col_idx] if col_idx < len(label_row) else None
            samples = []
            for row in data_rows:
                if col_idx < len(row) and row[col_idx] is not None:
                    samples.append(str(row[col_idx]).strip()[:60])
            columns.append({
                "index": col_idx,
                "header": str(header_val).strip() if header_val else f"Column {col_idx + 1}",
                "samples": samples[:3]
            })

        return {"columns": columns}

    def parse_tab_generic(self, sheet, header_row: int, mapping: dict) -> tuple:
        tasks = []
        flagged = []

        content_col  = mapping.get("content_col")
        deadline_col = mapping.get("deadline_col")
        staff_col    = mapping.get("staff_col")

        for row in sheet.iter_rows(min_row=header_row, values_only=True):
            if not any(row):
                continue

            content = self._cell_str(row[content_col]) if content_col is not None and content_col < len(row) else None
            if not content:
                continue

            deadline_raw = row[deadline_col] if deadline_col is not None and deadline_col < len(row) else None
            staff_text   = self._cell_str(row[staff_col]) if staff_col is not None and staff_col < len(row) else None

            is_recurring, recurrence_label = self.detect_recurring(str(deadline_raw) if deadline_raw else "")
            deadline = None if is_recurring else self.extract_deadline(deadline_raw)
            staff_names = self.extract_staff_names(staff_text)

            task = {
                "content": content,
                "deadline": deadline,
                "is_recurring": is_recurring,
                "recurrence_label": recurrence_label,
                "status": "pending",
                "staff_names": staff_names,
            }

            tasks.append(task)
            if not deadline and not is_recurring:
                flagged.append({**task, "flag_reason": "deadline not found"})

        return tasks, flagged

    def parse(self, tab_configs: list) -> dict:
        documents = []
        directives = []
        all_flagged = []
        all_staff = set()

        for cfg in tab_configs:
            sheet = self.wb[cfg["tab_name"]]
            tasks, flagged = self.parse_tab_generic(sheet, cfg["header_row"], cfg["mapping"])

            for task in tasks + flagged:
                for name in task.get("staff_names", []):
                    all_staff.add(name)

            if cfg.get("task_type") == "directive":
                directives.extend(tasks)
            else:
                documents.extend(tasks)

            all_flagged.extend(flagged)

        return {
            "documents": documents,
            "directives": directives,
            "flagged": all_flagged,
            "staff": list(all_staff),
            "summary": {
                "documents_parsed": len(documents),
                "directives_parsed": len(directives),
                "total_flagged": len(all_flagged),
                "staff_found": len(all_staff)
            }
        }
